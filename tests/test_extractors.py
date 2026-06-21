import docx as docx_lib
import pytest

from vectorizer.extractors import extract
from vectorizer.extractors import pdf as pdf_extractor


def test_extract_text(tmp_path):
    f = tmp_path / 'note.txt'
    f.write_text('hello world', encoding='utf-8')
    assert extract(f, 'text') == 'hello world'


def test_extract_md(tmp_path):
    f = tmp_path / 'note.md'
    f.write_text('# Title\n\nbody', encoding='utf-8')
    assert extract(f, 'md') == '# Title\n\nbody'


def test_extract_docx_headings_and_skips_empty(tmp_path):
    doc = docx_lib.Document()
    doc.add_paragraph('Big Title', style='Heading 1')
    doc.add_paragraph('Section', style='Heading 2')
    doc.add_paragraph('')          # blank -> skipped
    doc.add_paragraph('body text')
    f = tmp_path / 'doc.docx'
    doc.save(f)

    result = extract(f, 'docx')
    assert result == '# Big Title\n\n## Section\n\nbody text'


class _FakePage:
    def __init__(self, text):
        self._text = text

    def extract_text(self):
        return self._text


def test_extract_pdf_cleans_and_prefixes_pages(tmp_path, monkeypatch):
    pages = [
        _FakePage('first   line\n\n\n\nsecond  line'),
        _FakePage(''),               # empty page -> skipped
        _FakePage('third'),
    ]
    monkeypatch.setattr(pdf_extractor, 'PdfReader', lambda path: type('R', (), {'pages': pages})())

    f = tmp_path / 'doc.pdf'
    f.write_bytes(b'%PDF-1.4')
    result = extract(f, 'pdf')

    assert result == '[Page 1]\nfirst line\n\nsecond line\n\n[Page 3]\nthird'


def test_extract_pdf_handles_none_text(tmp_path, monkeypatch):
    monkeypatch.setattr(
        pdf_extractor, 'PdfReader',
        lambda path: type('R', (), {'pages': [_FakePage(None)]})(),
    )
    f = tmp_path / 'empty.pdf'
    f.write_bytes(b'%PDF-1.4')
    assert extract(f, 'pdf') == ''
